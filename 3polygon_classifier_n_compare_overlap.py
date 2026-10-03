"""
KMZ vs Excel Well Compare — Per Sumur
File 1 : KML/KMZ (acuan baris)
File 2 : Excel (urutan kolom: Nama, X DB, Y DB, X Evaluasi, Y Evaluasi,
         Jarak Deviasi, Status Verifikasi, Nama BKU)
Run    : streamlit run compare_kmz_excel.py
"""
import io
import os
import re
import zipfile
from collections import defaultdict
from math import radians, cos, sin, asin, sqrt
from xml.etree import ElementTree as ET

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from shapely.geometry import Point, Polygon
from shapely.ops import unary_union
from shapely.prepared import prep
try:
    from shapely.validation import make_valid
except ImportError:  # shapely < 1.8
    make_valid = lambda g: g.buffer(0)

# ── KONSTANTA ────────────────────────────────────────────────────────────────
EXCEL_COLS = ["Nama Sumur", "Koordinat Database X", "Koordinat Database Y",
              "Koordinat Hasil Evaluasi X", "Koordinat Hasil Evaluasi Y",
              "Jarak Deviasi (meter)", "Status Verifikasi", "Nama BKU"]

ST_SESUAI = "Sesuai Database"
ST_GESER = "Tidak Sesuai - Bergeser"
ST_TIDAK = "Tidak Ditemukan"


def status_priority(status):
    """0 = Sesuai Database, 1 = Bergeser, 2 = Tidak Ditemukan / lainnya."""
    s = str(status or "").strip().lower()
    if "sesuai database" in s and "tidak" not in s:
        return 0
    if "bergeser" in s:
        return 1
    return 2


def status_label(status):
    return {0: ST_SESUAI, 1: ST_GESER, 2: ST_TIDAK}[status_priority(status)]


# ── HELPERS ──────────────────────────────────────────────────────────────────
def haversine_m(lat1, lon1, lat2, lon2):
    R = 6_371_000
    p1, p2 = radians(lat1), radians(lat2)
    a = sin((p2 - p1) / 2) ** 2 + cos(p1) * cos(p2) * sin(radians(lon2 - lon1) / 2) ** 2
    return 2 * R * asin(sqrt(a))


def digits(name):
    return "".join(c for c in str(name) if c.isdigit())


def name_match(a, b):
    """True kalau nama identik (abaikan spasi/tanda baca/case) atau 4 digit terakhir sama."""
    na = re.sub(r"[^A-Z0-9]", "", str(a).upper())
    nb = re.sub(r"[^A-Z0-9]", "", str(b).upper())
    if na and na == nb:
        return True
    da, db = digits(a), digits(b)
    return len(da) >= 4 and len(db) >= 4 and da[-4:] == db[-4:]


def _tag(el):
    return el.tag.rsplit("}", 1)[-1]


# ── FILE 1: KML/KMZ ──────────────────────────────────────────────────────────
def read_kml_bytes(uploaded):
    raw = uploaded.getvalue()
    if uploaded.name.lower().endswith(".kmz"):
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            kmls = [n for n in z.namelist() if n.lower().endswith(".kml")]
            if not kmls:
                raise ValueError(f"Tidak ada .kml di dalam {uploaded.name}")
            kmls.sort(key=lambda n: (os.path.basename(n).lower() != "doc.kml", n))
            return z.read(kmls[0])
    return raw


def parse_kml_points(kml_bytes):
    """Namespace-agnostic. Ambil semua Placemark yang punya Point (termasuk dalam MultiGeometry)."""
    root = ET.fromstring(kml_bytes)
    recs = []
    for pm in root.iter():
        if _tag(pm) != "Placemark":
            continue
        name, ext = "", {}
        for ch in pm:
            if _tag(ch) == "name" and ch.text:
                name = ch.text.strip()
        for el in pm.iter():
            t = _tag(el)
            if t == "SimpleData":
                ext[el.get("name", "")] = (el.text or "").strip()
            elif t == "Data":
                for v in el:
                    if _tag(v) == "value":
                        ext[el.get("name", "")] = (v.text or "").strip()
        if not name:
            for k in ("NO_SUMUR", "Name_1", "name", "Name", "NAMA"):
                if ext.get(k):
                    name = ext[k]
                    break
        coord_text = None
        for el in pm.iter():
            if _tag(el) == "Point":
                for c in el:
                    if _tag(c) == "coordinates" and c.text:
                        coord_text = c.text
                break
        if not coord_text:
            continue
        parts = coord_text.strip().split()[0].split(",")
        try:
            lon, lat = float(parts[0]), float(parts[1])
        except (ValueError, IndexError):
            continue
        recs.append({"name": name, "lon": lon, "lat": lat,
                     "lon_str": parts[0].strip(), "lat_str": parts[1].strip(),
                     "order": len(recs)})
    return recs


# ── POLYGON ──────────────────────────────────────────────────────────────────
POLY_RULES = ["Hanya Info (tidak difilter)", "Lolos jika Dalam", "Lolos jika Luar"]


def _coords(text):
    pts = []
    for tok in (text or "").strip().split():
        p = tok.split(",")
        if len(p) >= 2:
            try:
                pts.append((float(p[0]), float(p[1])))
            except ValueError:
                pass
    return pts


def parse_kml_polygons(kml_bytes):
    """Namespace-agnostic. Ambil semua Polygon (outer + hole), auto-fix polygon invalid."""
    root = ET.fromstring(kml_bytes)
    polys = []
    for pg in root.iter():
        if _tag(pg) != "Polygon":
            continue
        outer, holes = None, []
        for b in pg:
            ring = None
            for el in b.iter():
                if _tag(el) == "coordinates":
                    ring = _coords(el.text)
                    break
            if not ring or len(ring) < 3:
                continue
            if _tag(b) == "outerBoundaryIs":
                outer = ring
            elif _tag(b) == "innerBoundaryIs":
                holes.append(ring)
        if outer:
            g = Polygon(outer, holes)
            polys.append(g if g.is_valid else make_valid(g))
    if not polys:  # fallback: LinearRing / LineString tertutup
        for el in root.iter():
            if _tag(el) in ("LinearRing", "LineString"):
                for c in el:
                    if _tag(c) == "coordinates":
                        ring = _coords(c.text)
                        if len(ring) >= 3:
                            g = Polygon(ring)
                            polys.append(g if g.is_valid else make_valid(g))
    return polys


def apply_polygons(df, polygons, lon_c, lat_c):
    """polygons: list dict {col, geom, rule}. Tambah kolom Dalam/Luar + Status Analisa Spasial."""
    if df.empty or not polygons:
        return df
    df = df.copy()
    for p in polygons:
        pg = prep(p["geom"])
        df[p["col"]] = [("Dalam" if pg.covers(Point(x, y)) else "Luar")
                        for x, y in zip(df[lon_c], df[lat_c])]

    def lolos(r):
        for p in polygons:
            if p["rule"] == "Lolos jika Dalam" and r[p["col"]] != "Dalam":
                return "Tidak Lolos"
            if p["rule"] == "Lolos jika Luar" and r[p["col"]] != "Luar":
                return "Tidak Lolos"
        return "Lolos"

    df["Status Analisa Spasial"] = df.apply(lolos, axis=1)
    return df


def sort_per_sumur(df, df_red):
    """
    Urutan: Lolos spasial dulu → Tidak Lolos; di dalamnya Berpasangan dulu → Tidak Berpasangan.
    Grup koordinat duplikat tetap berurutan. Kolom No dinomori ulang,
    referensi 'No Baris (Per Sumur)' di sheet Redundant ikut di-update.
    """
    if df.empty:
        return df, df_red
    df = df.copy()
    sp = df["Status Analisa Spasial"] if "Status Analisa Spasial" in df.columns else pd.Series("Lolos", index=df.index)
    df["_k_sp"] = (sp != "Lolos").astype(int)
    df["_k_pair"] = (df["Keterangan"] == "Tidak Berpasangan").astype(int)
    df["_k_row"] = range(len(df))
    # kunci grup = kunci terbaik anggota grup → anggota grup tidak terpisah
    df["_g_sp"] = df.groupby("_grup")["_k_sp"].transform("min")
    df["_g_pair"] = df.groupby("_grup")["_k_pair"].transform("min")
    df = df.sort_values(["_g_sp", "_g_pair", "_grup", "_k_row"], kind="stable").reset_index(drop=True)
    df["No"] = range(1, len(df) + 1)
    first_no = df.groupby("_grup")["No"].min()
    df = df.drop(columns=["_k_sp", "_k_pair", "_k_row", "_g_sp", "_g_pair"])
    if len(df_red):
        df_red = df_red.copy()
        df_red["No Baris (Per Sumur)"] = df_red["_grup"].map(first_no)
    return df, df_red


def sort_spasial(df):
    """Lolos spasial dulu, urutan asli dipertahankan."""
    if df.empty or "Status Analisa Spasial" not in df.columns:
        return df
    k = (df["Status Analisa Spasial"] != "Lolos").astype(int)
    return df.assign(_k=k).sort_values("_k", kind="stable").drop(columns="_k").reset_index(drop=True)


# ── FILE 2: EXCEL ────────────────────────────────────────────────────────────
def parse_excel(uploaded, sheet=0):
    df = pd.read_excel(uploaded, sheet_name=sheet, header=0)
    if df.shape[1] < 3:
        raise ValueError("Excel minimal 3 kolom (Nama, X, Y).")
    df = df.iloc[:, :8].copy()
    while df.shape[1] < 8:
        df[f"_pad{df.shape[1]}"] = None
    df.columns = EXCEL_COLS
    df = df.dropna(how="all").reset_index(drop=True)

    recs, skipped = [], []
    for i, r in df.iterrows():
        x = pd.to_numeric(r[EXCEL_COLS[1]], errors="coerce")
        y = pd.to_numeric(r[EXCEL_COLS[2]], errors="coerce")
        name = "(tanpa nama)" if pd.isna(r[EXCEL_COLS[0]]) else str(r[EXCEL_COLS[0]]).strip()
        info = {c: (None if pd.isna(r[c]) else r[c]) for c in EXCEL_COLS[3:]}
        rec = {"name": name, "lon": x, "lat": y, "order": i, **info,
               "_prio": status_priority(info["Status Verifikasi"])}
        if pd.isna(x) or pd.isna(y):
            skipped.append(rec)
            continue
        rec["lon_str"], rec["lat_str"] = repr(float(x)), repr(float(y))
        recs.append(rec)
    return recs, skipped, df


# ── CLUSTER & MATCH ──────────────────────────────────────────────────────────
def _cell(lat, lon, deg):
    return int(lat // deg), int(lon // deg)


def cluster_points(recs, dup_m):
    """dup_m = 0 → hanya koordinat identik. > 0 → gabung titik ≤ dup_m meter."""
    if dup_m <= 0:
        g = defaultdict(list)
        for r in recs:
            g[(r["lat"], r["lon"])].append(r)
        return [{"lat": v[0]["lat"], "lon": v[0]["lon"], "recs": v} for v in g.values()]
    deg = max(dup_m / 111_320.0, 1e-9)
    grid, clusters = defaultdict(list), []
    for r in recs:
        ci, cj = _cell(r["lat"], r["lon"], deg)
        hit = None
        for di in (-1, 0, 1):
            for dj in (-1, 0, 1):
                for k in grid.get((ci + di, cj + dj), []):
                    if haversine_m(r["lat"], r["lon"], clusters[k]["lat"], clusters[k]["lon"]) <= dup_m:
                        hit = k
                        break
                if hit is not None:
                    break
            if hit is not None:
                break
        if hit is None:
            clusters.append({"lat": r["lat"], "lon": r["lon"], "recs": [r]})
            grid[(ci, cj)].append(len(clusters) - 1)
        else:
            clusters[hit]["recs"].append(r)
    return clusters


def match_clusters(ca, cb, thr_m):
    """Greedy 1-to-1 antar lokasi, pasangan terdekat duluan."""
    deg = max(thr_m / 111_320.0, 1e-9)
    grid = defaultdict(list)
    for j, c in enumerate(cb):
        grid[_cell(c["lat"], c["lon"], deg)].append(j)
    cand = []
    for i, c in enumerate(ca):
        ci, cj = _cell(c["lat"], c["lon"], deg)
        for di in (-1, 0, 1):
            for dj in (-1, 0, 1):
                for j in grid.get((ci + di, cj + dj), []):
                    d = haversine_m(c["lat"], c["lon"], cb[j]["lat"], cb[j]["lon"])
                    if d <= thr_m:
                        cand.append((d, i, j))
    cand.sort()
    ua, ub, out = set(), set(), {}
    for d, i, j in cand:
        if i in ua or j in ub:
            continue
        ua.add(i); ub.add(j)
        out[i] = (j, d)
    return out


# ── PAIRING DALAM 1 GRUP ─────────────────────────────────────────────────────
def pair_group(kmz_recs, xl_recs):
    """
    kmz_recs : n nama (urutan file KMZ)
    xl_recs  : m nama Excel di koordinat pasangan
    Return   : list (kmz_rec, xl_rec, ket) panjang n, list redundant (sisa Excel)
    Rules:
      - Prioritas pasangan: Sesuai Database > Bergeser > Tidak Ditemukan.
      - m > n : n terbaik dipasang, sisa (Tidak Ditemukan duluan) → redundant.
      - m < n : semua Excel dipakai, baris sisa diisi ulang (berulang).
      - Di dalam yang terpilih, nama yang cocok (identik/4 digit akhir) dipasang duluan.
    """
    n, m = len(kmz_recs), len(xl_recs)
    ranked = sorted(xl_recs, key=lambda r: (r["_prio"], r["order"]))
    if m > n:
        chosen, redundant = ranked[:n], ranked[n:]
        redundant = sorted(redundant, key=lambda r: (-r["_prio"], r["order"]))
    else:
        chosen, redundant = ranked, []

    assign = [None] * n
    free = list(chosen)
    # 1) nama cocok
    for i, k in enumerate(kmz_recs):
        for x in free:
            if name_match(k["name"], x["name"]):
                assign[i] = (x, "Berpasangan")
                free.remove(x)
                break
    # 2) sisa terpilih sesuai urutan prioritas
    for i in range(n):
        if assign[i] is None and free:
            assign[i] = (free.pop(0), "Berpasangan")
    # 3) m < n → isi berulang, mulai dari prioritas terbaik
    rep = 0
    for i in range(n):
        if assign[i] is None:
            assign[i] = (ranked[rep % len(ranked)], "Pasangan Berulang")
            rep += 1
    return [(kmz_recs[i], assign[i][0], assign[i][1]) for i in range(n)], redundant


# ── BUILD ────────────────────────────────────────────────────────────────────
def build(kmz_recs, xl_recs, thr_m, dup_m, lbl1, lbl2):
    ck = cluster_points(kmz_recs, dup_m)
    cx = cluster_points(xl_recs, dup_m)
    mm = match_clusters(ck, cx, thr_m)

    order = sorted(range(len(ck)), key=lambda i: min(r["order"] for r in ck[i]["recs"]))
    rows, redundant_rows = [], []
    used_x, paired_x, red_x = set(), {}, {}
    no = 0
    for gid, ci in enumerate(order, start=1):
        members = sorted(ck[ci]["recs"], key=lambda r: r["order"])
        n = len(members)
        pair = mm.get(ci)
        if pair:
            j, dist = pair
            xs = cx[j]["recs"]
            pairs, red = pair_group(members, xs)
            for x in xs:
                used_x.add(id(x))
        else:
            dist, xs, red = None, [], []
            pairs = [(k, None, "Tidak Berpasangan") for k in members]

        for idx, (k, x, ket) in enumerate(pairs):
            no += 1
            if x is not None:
                paired_x.setdefault(id(x), (x, []))[1].append(no)
            row = {
                "No": no,
                f"Nama Sumur {lbl1}": k["name"],
                f"Longitude {lbl1}": k["lon"],
                f"Latitude {lbl1}": k["lat"],
                f"Jumlah Sumur di Koordinat ({lbl1})": n,
                f"Nama Sumur {lbl2}": x["name"] if x else "",
                "Koordinat Database X": x["lon"] if x else "",
                "Koordinat Database Y": x["lat"] if x else "",
                "Koordinat Hasil Evaluasi X": (x["Koordinat Hasil Evaluasi X"] if x else "") or "",
                "Koordinat Hasil Evaluasi Y": (x["Koordinat Hasil Evaluasi Y"] if x else "") or "",
                "Jarak Deviasi (meter)": (x["Jarak Deviasi (meter)"] if x else ""),
                "Status Verifikasi": (x["Status Verifikasi"] if x else "") or "",
                "Nama BKU": (x["Nama BKU"] if x else "") or "",
                f"Jumlah Sumur di Koordinat ({lbl2})": len(xs),
                f"Jarak {lbl1}-{lbl2} (m)": round(dist, 3) if dist is not None else "",
                "Kesamaan Nama": ("Sama" if name_match(k["name"], x["name"]) else "Berbeda") if x else "-",
                "Redundant Well": ", ".join(r["name"] for r in red) if (idx == 0 and red) else "",
                "Status Redundant Well": ", ".join(status_label(r["Status Verifikasi"]) for r in red)
                if (idx == 0 and red) else "",
                "Keterangan": ket,
                "_grup": gid, "_dup": n > 1,
            }
            if row["Jarak Deviasi (meter)"] is None:
                row["Jarak Deviasi (meter)"] = ""
            rows.append(row)

        for r in red:
            red_x[id(r)] = ", ".join(k["name"] for k in members)
            redundant_rows.append({
                f"Nama Sumur {lbl2}": r["name"],
                "Koordinat Database X": r["lon"], "Koordinat Database Y": r["lat"],
                "Status Verifikasi": r["Status Verifikasi"] or "",
                "Nama BKU": r["Nama BKU"] or "",
                f"Grup Sumur {lbl1}": ", ".join(k["name"] for k in members),
                "No Baris (Per Sumur)": rows[-n]["No"],
                "_grup": gid,
            })

    unmatched_x = [{
        f"Nama Sumur {lbl2}": r["name"],
        "Koordinat Database X": r["lon"], "Koordinat Database Y": r["lat"],
        "Koordinat Hasil Evaluasi X": r["Koordinat Hasil Evaluasi X"] or "",
        "Koordinat Hasil Evaluasi Y": r["Koordinat Hasil Evaluasi Y"] or "",
        "Jarak Deviasi (meter)": "" if r["Jarak Deviasi (meter)"] is None else r["Jarak Deviasi (meter)"],
        "Status Verifikasi": r["Status Verifikasi"] or "",
        "Nama BKU": r["Nama BKU"] or "",
    } for r in sorted(xl_recs, key=lambda r: r["order"]) if id(r) not in used_x]

    # semua sumur Excel + status pasangannya (untuk sheet Analisa Spasial File 2)
    nos_by_x = {i: v[1] for i, v in paired_x.items()}
    kmz_by_no = {r["No"]: r[f"Nama Sumur {lbl1}"] for r in rows}
    xl_all = []
    for r in sorted(xl_recs, key=lambda r: r["order"]):
        if id(r) in paired_x:
            stp, ref = "Berpasangan", ", ".join(kmz_by_no[n] for n in nos_by_x[id(r)])
        elif id(r) in red_x:
            stp, ref = "Redundant Well", red_x[id(r)]
        else:
            stp, ref = "Tidak Berpasangan", ""
        xl_all.append({
            f"Nama Sumur {lbl2}": r["name"],
            "Koordinat Database X": r["lon"], "Koordinat Database Y": r["lat"],
            "Koordinat Hasil Evaluasi X": r["Koordinat Hasil Evaluasi X"] or "",
            "Koordinat Hasil Evaluasi Y": r["Koordinat Hasil Evaluasi Y"] or "",
            "Jarak Deviasi (meter)": "" if r["Jarak Deviasi (meter)"] is None else r["Jarak Deviasi (meter)"],
            "Status Verifikasi": r["Status Verifikasi"] or "",
            "Nama BKU": r["Nama BKU"] or "",
            f"Status Pasangan ke {lbl1}": stp,
            f"Pasangan {lbl1}": ref,
            "_prio": r["_prio"], "_order": r["order"],
        })

    df_main = pd.DataFrame(rows)
    df_red = pd.DataFrame(redundant_rows)
    df_unx = pd.DataFrame(unmatched_x)
    df_xl_all = pd.DataFrame(xl_all)
    # paired: list (rec_excel, [No baris Per Sumur sebelum sort]) → dipakai filter Lolos
    return df_main, df_red, df_unx, ck, cx, list(paired_x.values()), df_xl_all


def bku_table(recs):
    """Rekap per Nama BKU: Sesuai, Bergeser, Tidak Ditemukan."""
    d = defaultdict(lambda: [0, 0, 0])
    for r in recs:
        bku = str(r.get("Nama BKU") or "(Tanpa Nama BKU)").strip()
        d[bku][r["_prio"]] += 1
    out = []
    for bku in sorted(d, key=str.lower):
        s, g, t = d[bku]
        out.append({"Nama BKU": bku, ST_SESUAI: s, ST_GESER: g,
                    "Total Ditemukan (Sesuai + Bergeser)": s + g, ST_TIDAK: t,
                    "Total Sumur": s + g + t})
    if out:
        tot = {"Nama BKU": "TOTAL"}
        for k in list(out[0].keys())[1:]:
            tot[k] = sum(o[k] for o in out)
        out.append(tot)
    return pd.DataFrame(out)


# ── EXCEL OUTPUT ─────────────────────────────────────────────────────────────
HDR_BLUE, HDR_GRAY, WHITE, ZEBRA = "1F4E79", "4472C4", "FFFFFF", "F2F2F2"
BLUE1, BLUE2 = "DDEBF7", "B4C6E7"
FILL_STATUS = {0: ("C6EFCE", "006100"), 1: ("FFEB9C", "9C5700"), 2: ("FFC7CE", "9C0006")}
THIN = Side(style="thin", color="AAAAAA")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
COORD_FMT = "0.##########"


def _fill(c):
    return PatternFill("solid", fgColor=c)


def write_table(ws, df, title, start_row=1, row_fills=None, coord_cols=()):
    vis = [c for c in df.columns if not str(c).startswith("_")]
    ncol = max(len(vis), 1)
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=ncol)
    t = ws.cell(row=start_row, column=1, value=title)
    t.font = Font(bold=True, size=12, color=WHITE)
    t.fill = _fill(HDR_BLUE)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[start_row].height = 22

    hr = start_row + 1
    for ci, col in enumerate(vis, 1):
        c = ws.cell(row=hr, column=ci, value=col)
        c.font = Font(bold=True, size=10, color=WHITE)
        c.fill = _fill(HDR_GRAY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[hr].height = 42

    for ri, (_, r) in enumerate(df.iterrows()):
        er = hr + 1 + ri
        base = (row_fills[ri] if row_fills and row_fills[ri] else (ZEBRA if ri % 2 else WHITE))
        is_total = str(r.get(vis[0], "")) == "TOTAL"
        for ci, col in enumerate(vis, 1):
            v = r[col]
            if v is None or (isinstance(v, float) and pd.isna(v)):
                v = ""
            c = ws.cell(row=er, column=ci, value=v)
            if col in coord_cols and isinstance(v, (int, float)) and v != "":
                c.number_format = COORD_FMT
            c.border = BORDER
            c.font = Font(size=9, bold=is_total)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.fill = _fill("D9E1F2" if is_total else base)
            if col in ("Status Verifikasi",) and v:
                bg, fg = FILL_STATUS[status_priority(v)]
                c.fill, c.font = _fill(bg), Font(size=9, bold=True, color=fg)
            elif col == "Keterangan" and v:
                bg, fg = {"Berpasangan": FILL_STATUS[0], "Pasangan Berulang": FILL_STATUS[1]}.get(v, FILL_STATUS[2])
                c.fill, c.font = _fill(bg), Font(size=9, bold=True, color=fg)
            elif col == "Status Analisa Spasial" and v:
                bg, fg = FILL_STATUS[0] if v == "Lolos" else FILL_STATUS[2]
                c.fill, c.font = _fill(bg), Font(size=9, bold=True, color=fg)
            elif str(col).startswith("Status Pasangan ke") and v:
                bg, fg = {"Berpasangan": FILL_STATUS[0], "Redundant Well": ("FCE4D6", "833C0B")}.get(v, FILL_STATUS[2])
                c.fill, c.font = _fill(bg), Font(size=9, bold=True, color=fg)
            elif col == "Redundant Well" and v:
                c.fill, c.font = _fill("FCE4D6"), Font(size=9, bold=True, color="833C0B")
    return hr + 1 + len(df)


def group_fills(df):
    fills, tog, prev = [], 0, None
    for _, r in df.iterrows():
        if r["_dup"]:
            if r["_grup"] != prev:
                tog ^= 1
            fills.append(BLUE1 if tog else BLUE2)
            prev = r["_grup"]
        else:
            fills.append(None)
            prev = None
    return fills


def set_widths(ws, df, default=16, wide=None):
    wide = wide or {}
    vis = [c for c in df.columns if not str(c).startswith("_")]
    for i, c in enumerate(vis, 1):
        ws.column_dimensions[get_column_letter(i)].width = wide.get(c, default)


def build_excel(df_main, df_red, df_unx, rekap_rows, df_bku_all, df_bku_pair, lbl1, lbl2,
                df_bku_lolos=None, df_xl_sp=None, df_bku_xl_lolos=None):
    wb = Workbook()
    coord_cols = {f"Longitude {lbl1}", f"Latitude {lbl1}", "Koordinat Database X",
                  "Koordinat Database Y", "Koordinat Hasil Evaluasi X", "Koordinat Hasil Evaluasi Y"}

    ws = wb.active
    ws.title = "Per Sumur"
    sort_note = ("Lolos spasial → Tidak Lolos, lalu Berpasangan → Tidak Berpasangan"
                 if "Status Analisa Spasial" in df_main.columns else "Berpasangan → Tidak Berpasangan")
    write_table(ws, df_main, f"PER SUMUR: 1 baris = 1 nama sumur {lbl1} | urutan: {sort_note} "
                f"| koordinat sama berurutan (biru)",
                row_fills=group_fills(df_main), coord_cols=coord_cols)
    set_widths(ws, df_main, wide={"No": 6, f"Nama Sumur {lbl1}": 20, f"Nama Sumur {lbl2}": 20,
                                  "Redundant Well": 30, "Status Redundant Well": 30, "Nama BKU": 22})
    ws.freeze_panes = "C3"

    ws = wb.create_sheet("Rekap")
    df_rek = pd.DataFrame(rekap_rows, columns=["Keterangan", "Jumlah"])
    r = write_table(ws, df_rek, "REKAP TOTAL")
    r = write_table(ws, df_bku_all, f"KLASIFIKASI STATUS PER NAMA BKU — SEMUA SUMUR {lbl2.upper()}",
                    start_row=r + 2)
    r = write_table(ws, df_bku_pair, f"KLASIFIKASI STATUS PER NAMA BKU — SUMUR {lbl2.upper()} "
                    f"YANG BERPASANGAN DENGAN {lbl1.upper()}", start_row=r + 2)
    if df_bku_lolos is not None:
        r = write_table(ws, df_bku_lolos, f"KLASIFIKASI STATUS PER NAMA BKU — SUMUR {lbl2.upper()} "
                        f"BERPASANGAN & LOLOS ANALISA SPASIAL (cek titik {lbl1})", start_row=r + 2)
    if df_bku_xl_lolos is not None:
        write_table(ws, df_bku_xl_lolos, f"KLASIFIKASI STATUS PER NAMA BKU — SEMUA SUMUR {lbl2.upper()} "
                    f"LOLOS ANALISA SPASIAL (cek koordinat Database {lbl2})", start_row=r + 2)
    ws.column_dimensions["A"].width = 52
    for col in "BCDEF":
        ws.column_dimensions[col].width = 20

    if "Status Analisa Spasial" in df_main.columns:
        df_l = df_main[df_main["Status Analisa Spasial"] == "Lolos"].reset_index(drop=True)
        if len(df_l):
            ws = wb.create_sheet("Lolos Analisa Spasial")
            write_table(ws, df_l, f"LOLOS ANALISA SPASIAL ({len(df_l)} baris)",
                        row_fills=group_fills(df_l), coord_cols=coord_cols)
            set_widths(ws, df_l, wide={"No": 6, f"Nama Sumur {lbl1}": 20, f"Nama Sumur {lbl2}": 20,
                                       "Redundant Well": 30, "Status Redundant Well": 30, "Nama BKU": 22})
            ws.freeze_panes = "C3"

    if len(df_red):
        ws = wb.create_sheet("Redundant Well")
        write_table(ws, df_red, f"REDUNDANT WELL: sisa nama {lbl2} yang tidak kebagian pasangan",
                    coord_cols=coord_cols)
        set_widths(ws, df_red, wide={f"Grup Sumur {lbl1}": 40, "Nama BKU": 22})

    if len(df_unx):
        ws = wb.create_sheet(f"{lbl2} Tdk Berpasangan"[:31])
        write_table(ws, df_unx, f"{lbl2.upper()} — TIDAK ADA PASANGAN DI {lbl1.upper()}"
                    + (" (Lolos spasial di atas)" if "Status Analisa Spasial" in df_unx.columns else ""),
                    coord_cols=coord_cols)
        set_widths(ws, df_unx, wide={"Nama BKU": 22})
        ws.freeze_panes = "B3"

    if df_xl_sp is not None and len(df_xl_sp):
        ws = wb.create_sheet(f"{lbl2} Analisa Spasial"[:31])
        n_l = int((df_xl_sp["Status Analisa Spasial"] == "Lolos").sum())
        write_table(ws, df_xl_sp, f"{lbl2.upper()} — ANALISA SPASIAL KOORDINAT DATABASE "
                    f"(Lolos {n_l} | Tidak Lolos {len(df_xl_sp) - n_l}) — Lolos di atas",
                    coord_cols=coord_cols)
        set_widths(ws, df_xl_sp, wide={"Nama BKU": 22, f"Pasangan {lbl1}": 30,
                                       f"Status Pasangan ke {lbl1}": 18})
        ws.freeze_panes = "B3"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_rekap(kmz_recs, xl_recs, xl_skipped, df_main, df_red, df_unx, ck, cx, lbl1, lbl2, thr, dup):
    ket = df_main["Keterangan"].value_counts() if len(df_main) else {}
    paired_names = len(xl_recs) - len(df_unx) - len(df_red)
    return [
        (f"Total nama sumur {lbl1}", len(kmz_recs)),
        (f"Total lokasi/koordinat unik {lbl1}", len(ck)),
        (f"Lokasi {lbl1} dengan ≥2 nama (duplikat)", sum(len(c['recs']) > 1 for c in ck)),
        (f"Total nama sumur {lbl2} (koordinat valid)", len(xl_recs)),
        (f"Baris {lbl2} dilewati (koordinat kosong/invalid)", len(xl_skipped)),
        (f"Total lokasi/koordinat unik {lbl2}", len(cx)),
        (f"Baris {lbl1} — Berpasangan", int(ket.get("Berpasangan", 0))),
        (f"Baris {lbl1} — Pasangan Berulang", int(ket.get("Pasangan Berulang", 0))),
        (f"Baris {lbl1} — Tidak Berpasangan", int(ket.get("Tidak Berpasangan", 0))),
        (f"Nama {lbl2} terpasang ke {lbl1}", paired_names),
        (f"Nama {lbl2} masuk Redundant Well", len(df_red)),
        (f"Nama {lbl2} tidak ada pasangan di {lbl1}", len(df_unx)),
    ] + spasial_rekap(df_main, lbl1) + [
        ("Threshold pasangan antar file (meter)", thr),
        ("Threshold duplikat dalam file (meter)", dup if dup > 0 else "0 (identik persis)"),
    ]


def spasial_rekap_xl(df, lbl2, lbl1):
    if df is None or "Status Analisa Spasial" not in df.columns:
        return []
    lol = df["Status Analisa Spasial"] == "Lolos"
    out = [(f"Nama {lbl2} — Lolos Analisa Spasial", int(lol.sum())),
           (f"Nama {lbl2} — Tidak Lolos Analisa Spasial", int((~lol).sum()))]
    stc = f"Status Pasangan ke {lbl1}"
    for stp in ("Berpasangan", "Redundant Well", "Tidak Berpasangan"):
        m = df[stc] == stp
        out.append((f"Nama {lbl2} — Lolos & {stp}", int((lol & m).sum())))
    return out


def spasial_rekap(df, lbl1):
    if "Status Analisa Spasial" not in df.columns:
        return []
    lol = df["Status Analisa Spasial"] == "Lolos"
    pair = df["Keterangan"] != "Tidak Berpasangan"
    return [
        (f"Baris {lbl1} — Lolos Analisa Spasial", int(lol.sum())),
        (f"Baris {lbl1} — Tidak Lolos Analisa Spasial", int((~lol).sum())),
        (f"Baris {lbl1} — Lolos & Berpasangan", int((lol & pair).sum())),
        (f"Baris {lbl1} — Lolos & Tidak Berpasangan", int((lol & ~pair).sum())),
        (f"Baris {lbl1} — Tidak Lolos & Berpasangan", int((~lol & pair).sum())),
        (f"Baris {lbl1} — Tidak Lolos & Tidak Berpasangan", int((~lol & ~pair).sum())),
    ]


# ── UI ───────────────────────────────────────────────────────────────────────
def label_of(f, default):
    return os.path.splitext(f.name)[0].replace("_", " ").replace("-", " ") if f else default


def main():
    st.set_page_config(page_title="KMZ vs Excel — Per Sumur", page_icon="🛢️", layout="wide")
    st.title("🛢️ KMZ vs Excel — Per Sumur")
    st.caption("Baris output = nama sumur File 1 (KMZ). Pasangan dari File 2 (Excel) "
               "berdasarkan jarak koordinat Database X/Y.")

    with st.sidebar:
        st.header("⚙️ Pengaturan")
        thr = st.number_input("Threshold pasangan antar file (m)", 0.01, 10_000.0, 1.0, 0.5,
                              help="Lokasi KMZ & Excel dianggap pasangan kalau jarak ≤ nilai ini.")
        dup = st.number_input("Threshold duplikat dalam file (m)", 0.0, 1_000.0, 0.0, 0.5,
                              help="0 = hanya koordinat identik persis dianggap 1 lokasi.")
        lbl1 = st.text_input("Label File 1", "KMZ")
        lbl2 = st.text_input("Label File 2", "Excel")
        st.markdown("---")
        use_poly = st.checkbox("🗺️ Gunakan Polygon Classifier", value=False)

    c1, c2 = st.columns(2)
    with c1:
        st.subheader("📂 File 1 — KML/KMZ")
        f1 = st.file_uploader("Upload KML/KMZ", type=["kml", "kmz"], key="f1")
    with c2:
        st.subheader("📂 File 2 — Excel")
        f2 = st.file_uploader("Upload Excel", type=["xlsx", "xls"], key="f2")
        sheet = 0
        if f2:
            xls = pd.ExcelFile(f2)
            sheet = st.selectbox("Sheet", xls.sheet_names)

    poly_slots = []
    if use_poly:
        st.markdown("---")
        st.subheader("🗺️ Polygon Classifier (1–5 polygon)")
        st.caption("Per Sumur dicek pakai koordinat File 1 (KMZ). Sheet File 2 dicek pakai "
                   "Koordinat Database X/Y. Baris Lolos ditaruh paling atas.")
        for i in range(1, 6):
            with st.expander(f"Polygon {i}" + (" (wajib)" if i == 1 else " (opsional)"), expanded=(i == 1)):
                pf = st.file_uploader(f"File Polygon {i}", type=["kml", "kmz"], key=f"poly_{i}",
                                      label_visibility="collapsed")
                if pf:
                    pn = st.text_input(f"Nama Polygon {i}", os.path.splitext(pf.name)[0], key=f"pn_{i}")
                    pr = st.selectbox(f"Rule Polygon {i}", POLY_RULES, index=1, key=f"pr_{i}")
                    poly_slots.append({"idx": i, "file": pf, "name": pn, "rule": pr})

    if not st.button("🚀 PROSES", type="primary", use_container_width=True):
        return
    if not f1 or not f2:
        st.error("Upload 2 file dulu.")
        return
    if use_poly and not poly_slots:
        st.error("Polygon Classifier aktif, upload minimal 1 polygon.")
        return

    try:
        with st.spinner("Memproses..."):
            kmz = parse_kml_points(read_kml_bytes(f1))
            xl, xl_skip, _ = parse_excel(f2, sheet)
            if not kmz:
                st.error("Tidak ada titik terbaca di KMZ.")
                return
            if not xl:
                st.error("Tidak ada baris valid di Excel.")
                return

            df_main, df_red, df_unx, ck, cx, paired, df_xl_all = build(kmz, xl, thr, dup, lbl1, lbl2)
            xl_paired = [x for x, _ in paired]

            polygons = []
            for sl in poly_slots:
                geoms = parse_kml_polygons(read_kml_bytes(sl["file"]))
                if not geoms:
                    st.warning(f"Polygon {sl['idx']}: tidak ada polygon terbaca, dilewati.")
                    continue
                polygons.append({"col": f"Polygon {sl['idx']} ({sl['name']})",
                                 "geom": unary_union(geoms), "rule": sl["rule"]})
            df_main = apply_polygons(df_main, polygons, f"Longitude {lbl1}", f"Latitude {lbl1}")
            df_unx = sort_spasial(apply_polygons(df_unx, polygons,
                                                 "Koordinat Database X", "Koordinat Database Y"))
            df_xl_sp, df_bku_xl_lolos = None, None
            if polygons:
                df_xl_sp = sort_spasial(apply_polygons(df_xl_all, polygons,
                                                       "Koordinat Database X", "Koordinat Database Y"))
                lol_mask = df_xl_sp["Status Analisa Spasial"] == "Lolos"
                df_bku_xl_lolos = bku_table([{"Nama BKU": b or None, "_prio": p} for b, p in
                                             zip(df_xl_sp.loc[lol_mask, "Nama BKU"],
                                                 df_xl_sp.loc[lol_mask, "_prio"])])

            df_bku_lolos = None
            if polygons:
                lolos_no = set(df_main.loc[df_main["Status Analisa Spasial"] == "Lolos", "No"])
                df_bku_lolos = bku_table([x for x, nos in paired if lolos_no.intersection(nos)])

            df_main, df_red = sort_per_sumur(df_main, df_red)

            rekap = make_rekap(kmz, xl, xl_skip, df_main, df_red, df_unx, ck, cx,
                               lbl1, lbl2, thr, dup)
            rekap = rekap[:-2] + spasial_rekap_xl(df_xl_sp, lbl2, lbl1) + rekap[-2:]
            df_bku_all = bku_table(xl + xl_skip)
            df_bku_pair = bku_table(xl_paired)
            out = build_excel(df_main, df_red, df_unx, rekap, df_bku_all, df_bku_pair, lbl1, lbl2,
                              df_bku_lolos, df_xl_sp, df_bku_xl_lolos)

        m = st.columns(6 if polygons else 5)
        m[0].metric(f"Nama {lbl1}", len(kmz))
        m[1].metric(f"Nama {lbl2}", len(xl))
        m[2].metric("Berpasangan", int((df_main["Keterangan"] == "Berpasangan").sum()))
        m[3].metric("Tidak Berpasangan", int((df_main["Keterangan"] == "Tidak Berpasangan").sum()))
        m[4].metric("Redundant Well", len(df_red))
        if polygons:
            m[5].metric("Lolos Spasial", int((df_main["Status Analisa Spasial"] == "Lolos").sum()))
        if xl_skip:
            st.warning(f"{len(xl_skip)} baris Excel dilewati (koordinat Database kosong/invalid).")

        vis = lambda d: d[[c for c in d.columns if not str(c).startswith("_")]]
        tab_names = [f"Per Sumur ({len(df_main)})", "Rekap",
                     f"Redundant ({len(df_red)})", f"{lbl2} Tdk Berpasangan ({len(df_unx)})"]
        if df_xl_sp is not None:
            tab_names.append(f"{lbl2} Analisa Spasial ({len(df_xl_sp)})")
        tabs = st.tabs(tab_names)
        with tabs[0]:
            st.dataframe(vis(df_main), use_container_width=True, height=450)
        with tabs[1]:
            st.dataframe(pd.DataFrame(rekap, columns=["Keterangan", "Jumlah"]).astype(str),
                         use_container_width=True, hide_index=True)
            st.markdown("**Status per Nama BKU — semua sumur File 2**")
            st.dataframe(df_bku_all, use_container_width=True, hide_index=True)
            st.markdown("**Status per Nama BKU — sumur File 2 yang berpasangan**")
            st.dataframe(df_bku_pair, use_container_width=True, hide_index=True)
            if df_bku_lolos is not None:
                st.markdown("**Status per Nama BKU — berpasangan & lolos spasial**")
                st.dataframe(df_bku_lolos, use_container_width=True, hide_index=True)
            if df_bku_xl_lolos is not None:
                st.markdown("**Status per Nama BKU — semua sumur File 2 lolos spasial**")
                st.dataframe(df_bku_xl_lolos, use_container_width=True, hide_index=True)
        with tabs[2]:
            st.dataframe(vis(df_red), use_container_width=True) if len(df_red) else st.info("Kosong")
        with tabs[3]:
            st.dataframe(df_unx, use_container_width=True) if len(df_unx) else st.info("Kosong")
        if df_xl_sp is not None:
            with tabs[4]:
                st.dataframe(vis(df_xl_sp), use_container_width=True, height=450)

        st.download_button("📥 Download Excel", out,
                           file_name=f"Compare_{lbl1}_vs_{lbl2}.xlsx".replace(" ", "_"),
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           type="primary", use_container_width=True)
    except Exception as e:
        import traceback
        st.error(f"Error: {e}")
        st.code(traceback.format_exc())


if __name__ == "__main__":
    main()
